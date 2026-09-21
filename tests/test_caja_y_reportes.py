from datetime import date, datetime, time, timedelta
from decimal import Decimal

import pytest

from app.database import db
from app.models import (AsignacionTurno, Cajero, Caja, DetalleVenta, RegistroUso, Turno,
                        Venta)
from app.services import cuadre, reportes, turnos
from app.services.errores import ErrorNegocio
from tests.conftest import (ADMIN, CAJERO_REAL, CAJERO_REAL_MEDIO, abrir_caja, unidad,
                            usuario, vencer_todo, cerrar_caja)


# ------------------------------------------------------------------ caja y minutos
def test_no_se_vende_sin_caja_abierta(app, como):
    cajero = como(CAJERO_REAL)
    cajero.post('/caja/venta', data={'id_flota': unidad('Real Plaza').id_flota, 'bloques': 1, 'metodo_pago': 'efectivo'})
    assert Venta.query.count() == 0


def test_cierre_de_caja_cuadra_efectivo_y_minutos(app, como):
    cajero = como(CAJERO_REAL)
    abrir_caja(cajero, 50)
    u = unidad('Real Plaza')
    cajero.post('/caja/venta', data={'id_flota': u.id_flota, 'bloques': 2, 'metodo_pago': 'efectivo'})   # 20 efectivo, 24 min
    cajero.post('/caja/venta', data={'id_flota': unidad('Real Plaza', 3).id_flota, 'bloques': 1, 'metodo_pago': 'yape'})  # 10 yape, 12 min
    cajero.post('/caja/movimiento', data={'tipo': 'egreso', 'monto': '5', 'descripcion': 'Agua'})

    # Con los vehiculos todavia en uso no se puede cerrar
    cerrar_caja(cajero, '60')
    assert Caja.query.one().estado == 'abierta'

    vencer_todo()                                                   # el tiempo termina solo
    cerrar_caja(cajero, '60')       # 50 + 20 - 5 = 65 esperado; faltan 5
    db.session.expire_all()
    caja = Caja.query.one()
    assert caja.estado == 'cerrada'
    assert caja.efectivo_esperado == Decimal('65.00') and caja.diferencia == Decimal('-5.00')
    assert (caja.total_efectivo, caja.total_yape) == (Decimal('20.00'), Decimal('10.00'))
    assert (caja.minutos_vendidos, caja.minutos_usados) == (36, 36)


def test_finalizar_antes_registra_minutos_reales_y_alerta_diferencias(app, como):
    cajero = como(CAJERO_REAL)
    abrir_caja(cajero)
    cajero.post('/caja/venta', data={'id_flota': unidad('Real Plaza').id_flota, 'bloques': 1, 'metodo_pago': 'efectivo'})
    uso = RegistroUso.query.one()
    assert uso.estado == 'en_curso'
    uso.hora_inicio = datetime.now() - timedelta(minutes=5)         # llevaba 5 de 12 minutos
    db.session.commit()

    assert cajero.post(f'/caja/uso/{uso.id_uso}/finalizar', data={'motivo': ''}).status_code == 302
    db.session.expire_all()
    assert RegistroUso.query.one().estado == 'en_curso'             # sin motivo no finaliza

    cajero.post(f'/caja/uso/{uso.id_uso}/finalizar', data={'motivo': 'El nino se asusto'})
    db.session.expire_all()
    uso = RegistroUso.query.one()
    assert (uso.estado, uso.minutos_usados, uso.finalizado_por, uso.motivo_fin) == ('finalizado', 5, 'cajero', 'El nino se asusto')

    fila = next(f for f in cuadre.cuadre_diario(date.today())['filas'] if f['cajero'] == 'Cajero 1 Real Plaza')
    assert (fila['minutos_vendidos'], fila['minutos_usados'], fila['diferencia_minutos']) == (12, 5, 7)
    assert any('7 min m' in a for a in fila['alertas'])             # alerta para el administrador


def test_sin_diferencias_no_hay_alerta(app, como):
    cajero = como(CAJERO_REAL)
    abrir_caja(cajero)
    cajero.post('/caja/venta', data={'id_flota': unidad('Real Plaza').id_flota, 'bloques': 1, 'metodo_pago': 'efectivo'})
    vencer_todo()
    cerrar_caja(cajero, '60')       # 50 + 10
    db.session.expire_all()
    fila = next(f for f in cuadre.cuadre_diario(date.today())['filas'] if f['cajero'] == 'Cajero 1 Real Plaza')
    assert fila['alertas'] == [] and fila['estado_caja'] == 'cerrada'


# ------------------------------------------------------------------------ reportes
def _semana_pasada():
    lunes = turnos.lunes_de(date.today()) - timedelta(days=7)
    return lunes, lunes + timedelta(days=6)


def _preparar_semana():
    """Completo: 5 turnos de 8 h = 40 h y S/ 400. Medio: 5 turnos de 4 h = 20 h y S/ 100."""
    lunes, domingo = _semana_pasada()
    AsignacionTurno.query.delete()
    completo = Turno.query.filter_by(nombre='Completo mañana').one()
    medio = Turno.query.filter_by(nombre='Medio mañana').one()
    c_full = Cajero.query.filter_by(id_usuario=usuario(CAJERO_REAL).id_usuario).one()
    c_half = Cajero.query.filter_by(id_usuario=usuario(CAJERO_REAL_MEDIO).id_usuario).one()
    u = unidad('Real Plaza')
    for i in range(5):
        dia = lunes + timedelta(days=i)
        db.session.add(AsignacionTurno(id_cajero=c_full.id_cajero, id_turno=completo.id_turno, fecha=dia))
        db.session.add(AsignacionTurno(id_cajero=c_half.id_cajero, id_turno=medio.id_turno, fecha=dia))
        for cajero, monto, metodo in ((c_full, 80, 'efectivo'), (c_half, 20, 'yape')):
            v = Venta(fecha=datetime.combine(dia, time(13, 0)), total=Decimal(monto), tipo_pago=metodo,
                      estado='pagada', origen='presencial', id_usuario=cajero.id_usuario,
                      id_sucursal=cajero.id_sucursal)
            v.detalles.append(DetalleVenta(id_flota=u.id_flota, precio=Decimal(10), cantidad=monto // 10,
                                           minutos=monto // 10 * 12, subtotal=Decimal(monto), fecha_servicio=dia))
            db.session.add(v)
    db.session.commit()
    return lunes, domingo


def test_reporte_por_cajero_con_horas_y_ventas_por_hora(app):
    lunes, domingo = _preparar_semana()
    r = reportes.generar(lunes, domingo)
    por = {c['cajero']: c for c in r['por_cajero']}
    full, half = por['Cajero 1 Real Plaza'], por['Cajero 2 Real Plaza']

    assert (full['total'], full['ventas'], full['minutos']) == (400.0, 5, 480)
    assert (half['total'], half['ventas'], half['minutos']) == (100.0, 5, 120)
    assert (full['horas'], half['horas']) == (40.0, 20.0)                 # 8 h vs 4 h por turno
    assert (full['ventas_por_hora'], half['ventas_por_hora']) == (10.0, 5.0)  # comparación justa
    assert full['ticket_promedio'] == 80.0 and half['ticket_promedio'] == 20.0
    assert r['resumen'] == {'total': 500.0, 'ventas': 10, 'minutos': 600, 'ticket_promedio': 50.0}


def test_reporte_por_sede_tipo_y_metodo(app):
    lunes, domingo = _preparar_semana()
    r = reportes.generar(lunes, domingo)
    assert {s['sede']: s['total'] for s in r['por_sede']}['Real Plaza'] == 500.0
    assert {s['sede']: s['total'] for s in r['por_sede']}['Mega Plaza'] == 0.0
    assert r['por_tipo'][0]['total'] == 500.0
    assert {m['metodo']: m['total'] for m in r['por_metodo']} == {'Efectivo': 400.0, 'Yape': 100.0}
    assert sum(d['total'] for d in r['por_dia']) == 500.0 and len(r['por_dia']) == 7


def test_filtro_por_sede(app):
    lunes, domingo = _preparar_semana()
    from tests.conftest import sede
    assert reportes.generar(lunes, domingo, sede('Mega Plaza').id_sucursal)['resumen']['total'] == 0
    assert reportes.generar(lunes, domingo, sede('Real Plaza').id_sucursal)['resumen']['total'] == 500.0


def test_rangos_semanal_mensual_y_personalizado(app):
    hoy = date(2026, 9, 21)                                   # lunes
    assert reportes.resolver_rango('semanal', hoy=hoy)[:2] == (date(2026, 9, 21), date(2026, 9, 27))
    assert reportes.resolver_rango('semanal', referencia='2026-09-23')[:2] == (date(2026, 9, 21), date(2026, 9, 27))
    assert reportes.resolver_rango('mensual', hoy=hoy)[:2] == (date(2026, 9, 1), date(2026, 9, 30))
    assert reportes.resolver_rango('mensual', referencia='2026-02-10')[:2] == (date(2026, 2, 1), date(2026, 2, 28))
    assert reportes.resolver_rango('personalizado', '2026-09-05', '2026-09-12')[:2] == (date(2026, 9, 5), date(2026, 9, 12))
    with pytest.raises(ErrorNegocio):
        reportes.resolver_rango('personalizado', '2026-09-12', '2026-09-05')
    with pytest.raises(ErrorNegocio):
        reportes.resolver_rango('personalizado', 'ayer', 'hoy')


def test_horas_estimadas_si_no_hay_turnos_asignados(app):
    lunes, domingo = _semana_pasada()
    AsignacionTurno.query.delete()
    u = unidad('Real Plaza')
    cajero = Cajero.query.filter_by(id_usuario=usuario(CAJERO_REAL).id_usuario).one()
    for i in range(3):        # 3 días con ventas -> 3 x 8 h
        v = Venta(fecha=datetime.combine(lunes + timedelta(days=i), time(12, 0)), total=Decimal(10), tipo_pago='efectivo',
                  estado='pagada', id_usuario=cajero.id_usuario, id_sucursal=cajero.id_sucursal)
        v.detalles.append(DetalleVenta(id_flota=u.id_flota, precio=Decimal(10), cantidad=1, minutos=12,
                                       subtotal=Decimal(10), fecha_servicio=lunes))
        db.session.add(v)
    db.session.commit()
    fila = next(c for c in reportes.por_cajero(lunes, domingo) if c['cajero'] == 'Cajero 1 Real Plaza')
    assert (fila['horas'], fila['fuente_horas']) == (24.0, 'estimado')


def test_exportacion_a_excel_y_pdf(app, como):
    _preparar_semana()
    admin = como(ADMIN)
    lunes, domingo = _semana_pasada()
    filtro = f'?rango=personalizado&desde={lunes}&hasta={domingo}'
    excel = admin.get('/reportes/exportar/excel' + filtro)
    assert excel.status_code == 200 and excel.data[:2] == b'PK'         # .xlsx es un zip
    pdf = admin.get('/reportes/exportar/pdf' + filtro)
    assert pdf.status_code == 200 and pdf.data[:4] == b'%PDF'
    assert admin.get('/reportes/exportar/doc').status_code == 404

    from openpyxl import load_workbook
    import io
    libro = load_workbook(io.BytesIO(excel.data))
    assert 'Ventas por cajero' in libro.sheetnames


def test_paginas_de_reporte_se_muestran(app, como):
    admin = como(ADMIN)
    for ruta in ('/reportes/dashboard', '/reportes/', '/reportes/cuadre', '/reportes/dashboard?rango=mensual',
                 '/admin/turnos', '/admin/usuarios', '/admin/flota', '/admin/tarifas', '/admin/sedes',
                 '/admin/mantenimiento/'):
        assert admin.get(ruta).status_code == 200, ruta
    assert admin.get('/reportes/?rango=personalizado&desde=2026-09-10&hasta=2026-09-01').status_code == 400


# ------------------------------------------------------------------------- turnos
def test_asignacion_semanal_y_validacion_de_tipo(app, como):
    admin = como(ADMIN)
    lunes = turnos.lunes_de(date.today()) + timedelta(days=14)
    cajero = Cajero.query.filter_by(id_usuario=usuario(CAJERO_REAL_MEDIO).id_usuario).one()   # medio tiempo
    medio = Turno.query.filter_by(nombre='Medio tarde').one()
    completo = Turno.query.filter_by(nombre='Completo tarde').one()

    r = admin.post('/admin/turnos/guardar', data={'semana': lunes.isoformat(),
                                                  f't_{cajero.id_cajero}_{lunes}': medio.id_turno})
    assert r.status_code == 302
    assert AsignacionTurno.query.filter_by(id_cajero=cajero.id_cajero, fecha=lunes).one().id_turno == medio.id_turno

    admin.post('/admin/turnos/guardar', data={'semana': lunes.isoformat(),
                                              f't_{cajero.id_cajero}_{lunes}': completo.id_turno})
    db.session.expire_all()      # un cajero de medio tiempo no puede recibir un turno completo
    assert AsignacionTurno.query.filter_by(id_cajero=cajero.id_cajero, fecha=lunes).one().id_turno == medio.id_turno


def test_importar_horario_desde_csv(app, como):
    import io
    admin = como(ADMIN)
    lunes = turnos.lunes_de(date.today()) + timedelta(days=21)
    csv = ('correo,Lunes,Martes,Miércoles,Jueves,Viernes,Sábado,Domingo\n'
           f'{CAJERO_REAL},Completo mañana,Completo mañana,Completo tarde,Completo tarde,libre,Completo mañana,libre\n')
    r = admin.post('/admin/turnos/importar', data={'semana': lunes.isoformat(),
                   'archivo': (io.BytesIO(csv.encode('utf-8')), 'horario.csv')})
    assert r.status_code == 302
    cajero = Cajero.query.filter_by(id_usuario=usuario(CAJERO_REAL).id_usuario).one()
    asignados = AsignacionTurno.query.filter(AsignacionTurno.id_cajero == cajero.id_cajero,
                                             AsignacionTurno.fecha >= lunes,
                                             AsignacionTurno.fecha < lunes + timedelta(days=7)).count()
    assert asignados == 5
