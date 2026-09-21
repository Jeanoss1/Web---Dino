"""Reportes de ventas para el administrador y su exportación a Excel / PDF."""
import calendar
import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from sqlalchemy import func

from app.database import db
from app.models.Cajero import Cajero
from app.models.DetalleVenta import DetalleVenta
from app.models.Flota import Flota
from app.models.Sucursal import Sucursal
from app.models.usuario import Usuario
from app.models.Venta import Venta
from app.services import turnos
from app.services import validadores as v
from app.services.errores import ErrorNegocio

ETIQUETA_TURNO = {'completo': 'Tiempo completo (8 h)', 'medio': 'Medio tiempo (4 h)'}


def _f(valor):
    return float(valor or 0)


# --------------------------------------------------------------------- rangos
def resolver_rango(rango, desde=None, hasta=None, referencia=None, hoy=None):
    """Devuelve (desde, hasta, etiqueta) para semanal, mensual o personalizado."""
    hoy = hoy or date.today()
    if rango == 'personalizado':
        d, h = v.fecha(desde, 'fecha inicial'), v.fecha(hasta, 'fecha final')
        if d > h:
            raise ErrorNegocio('La fecha inicial no puede ser posterior a la final.')
        if (h - d).days > 366:
            raise ErrorNegocio('El rango máximo es de un año.')
        return d, h, f'Del {d:%d/%m/%Y} al {h:%d/%m/%Y}'
    ref = v.fecha(referencia, 'fecha de referencia') if referencia else hoy
    if rango == 'diario':
        return ref, ref, f'Día {ref:%d/%m/%Y}'
    if rango == 'mensual':
        d = ref.replace(day=1)
        h = ref.replace(day=calendar.monthrange(ref.year, ref.month)[1])
        return d, h, f'Mes de {d:%m/%Y}'
    d = turnos.lunes_de(ref)                     # semanal (por defecto)
    h = d + timedelta(days=6)
    return d, h, f'Semana del {d:%d/%m/%Y} al {h:%d/%m/%Y}'


def _limites(desde, hasta):
    return (datetime.combine(desde, datetime.min.time()),
            datetime.combine(hasta + timedelta(days=1), datetime.min.time()))


def _filtro(desde, hasta, id_sucursal):
    ini, fin = _limites(desde, hasta)
    cond = [Venta.estado == 'pagada', Venta.fecha >= ini, Venta.fecha < fin]
    if id_sucursal:
        cond.append(Venta.id_sucursal == id_sucursal)
    return cond


# ------------------------------------------------------------------- consultas
def por_cajero(desde, hasta, id_sucursal=None):
    cond = _filtro(desde, hasta, id_sucursal)
    totales = {u: (t, n) for u, t, n in db.session.query(
        Venta.id_usuario, func.sum(Venta.total), func.count(Venta.id_venta)
    ).filter(*cond).group_by(Venta.id_usuario)}
    minutos = dict(db.session.query(Venta.id_usuario, func.sum(DetalleVenta.minutos))
                   .join(DetalleVenta, DetalleVenta.id_venta == Venta.id_venta)
                   .filter(*cond).group_by(Venta.id_usuario).all())

    filas, vistos = [], set()
    for cajero in turnos.cajeros_activos(id_sucursal):
        vistos.add(cajero.id_usuario)
        total, n = totales.get(cajero.id_usuario, (0, 0))
        horas, fuente = turnos.horas_trabajadas(cajero, desde, hasta)
        filas.append(_fila(cajero.id_usuario, cajero.nombre, cajero.sucursal.nombre,
                           ETIQUETA_TURNO.get(cajero.tipo_turno, cajero.tipo_turno),
                           total, n, minutos.get(cajero.id_usuario, 0), horas, fuente))
    # Ventas registradas por quien ya no es cajero activo o por un administrador
    for id_usuario, (total, n) in totales.items():
        if id_usuario in vistos:
            continue
        u = db.session.get(Usuario, id_usuario)
        sede = '—'
        if u.cajero:
            sede = u.cajero.sucursal.nombre
        filas.append(_fila(id_usuario, u.nombre_completo, sede, u.rol_nombre.capitalize(),
                           total, n, minutos.get(id_usuario, 0), 0, 'n/a'))
    return sorted(filas, key=lambda f: f['total'], reverse=True)


def _fila(id_usuario, nombre, sede, turno, total, n, minutos, horas, fuente):
    total = _f(total)
    return {
        'id_usuario': id_usuario, 'cajero': nombre, 'sede': sede, 'turno': turno,
        'total': total, 'ventas': int(n or 0), 'minutos': int(minutos or 0),
        'ticket_promedio': round(total / n, 2) if n else 0.0,
        'horas': round(horas, 1), 'fuente_horas': fuente,
        'ventas_por_hora': round(total / horas, 2) if horas else 0.0,
    }


def por_sede(desde, hasta, id_sucursal=None):
    cond = _filtro(desde, hasta, id_sucursal)
    datos = {s: (t, n) for s, t, n in db.session.query(
        Venta.id_sucursal, func.sum(Venta.total), func.count(Venta.id_venta)
    ).filter(*cond).group_by(Venta.id_sucursal)}
    minutos = dict(db.session.query(Venta.id_sucursal, func.sum(DetalleVenta.minutos))
                   .join(DetalleVenta, DetalleVenta.id_venta == Venta.id_venta)
                   .filter(*cond).group_by(Venta.id_sucursal).all())
    sedes = Sucursal.query.order_by(Sucursal.nombre)
    if id_sucursal:
        sedes = sedes.filter_by(id_sucursal=id_sucursal)
    return [{'sede': s.nombre, 'total': _f(datos.get(s.id_sucursal, (0, 0))[0]),
             'ventas': int(datos.get(s.id_sucursal, (0, 0))[1] or 0),
             'minutos': int(minutos.get(s.id_sucursal, 0) or 0)}
            for s in sedes if s.id_sucursal in datos or s.estado == 'activa']


def por_tipo(desde, hasta, id_sucursal=None):
    cond = _filtro(desde, hasta, id_sucursal)
    filas = (db.session.query(Flota.tipo, func.sum(DetalleVenta.subtotal),
                              func.sum(DetalleVenta.minutos), func.count(DetalleVenta.id_detalle_venta))
             .join(DetalleVenta, DetalleVenta.id_flota == Flota.id_flota)
             .join(Venta, Venta.id_venta == DetalleVenta.id_venta)
             .filter(*cond).group_by(Flota.tipo).all())
    return sorted(({'tipo': t.capitalize(), 'total': _f(s), 'minutos': int(m or 0), 'alquileres': int(n)}
                   for t, s, m, n in filas), key=lambda x: x['total'], reverse=True)


def por_metodo(desde, hasta, id_sucursal=None):
    cond = _filtro(desde, hasta, id_sucursal)
    filas = (db.session.query(Venta.tipo_pago, func.sum(Venta.total), func.count(Venta.id_venta))
             .filter(*cond).group_by(Venta.tipo_pago).all())
    return [{'metodo': (m or '—').capitalize(), 'total': _f(t), 'ventas': int(n)} for m, t, n in filas]


def por_dia(desde, hasta, id_sucursal=None):
    cond = _filtro(desde, hasta, id_sucursal)
    acumulado = {}
    for fecha, total in db.session.query(Venta.fecha, Venta.total).filter(*cond):
        acumulado[fecha.date()] = acumulado.get(fecha.date(), Decimal('0')) + total
    dias = [desde + timedelta(days=i) for i in range((hasta - desde).days + 1)]
    return [{'fecha': d.isoformat(), 'total': _f(acumulado.get(d))} for d in dias]


def generar(desde, hasta, id_sucursal=None, etiqueta=''):
    cajeros = por_cajero(desde, hasta, id_sucursal)
    total = sum(c['total'] for c in cajeros)
    ventas = sum(c['ventas'] for c in cajeros)
    return {
        'desde': desde, 'hasta': hasta, 'etiqueta': etiqueta, 'id_sucursal': id_sucursal,
        'resumen': {
            'total': round(total, 2), 'ventas': ventas,
            'minutos': sum(c['minutos'] for c in cajeros),
            'ticket_promedio': round(total / ventas, 2) if ventas else 0.0,
        },
        'por_cajero': cajeros,
        'por_sede': por_sede(desde, hasta, id_sucursal),
        'por_tipo': por_tipo(desde, hasta, id_sucursal),
        'por_metodo': por_metodo(desde, hasta, id_sucursal),
        'por_dia': por_dia(desde, hasta, id_sucursal),
    }


# ------------------------------------------------------------------ exportación
def _tablas(r):
    """(título, encabezados, filas) de cada bloque del reporte."""
    return [
        ('Resumen', ['Indicador', 'Valor'], [
            ['Total vendido (S/)', r['resumen']['total']], ['Número de ventas', r['resumen']['ventas']],
            ['Minutos vendidos', r['resumen']['minutos']],
            ['Ticket promedio (S/)', r['resumen']['ticket_promedio']]]),
        ('Ventas por cajero',
         ['Cajero', 'Sede', 'Turno', 'Total (S/)', 'Ventas', 'Minutos', 'Ticket prom. (S/)',
          'Horas trabajadas', 'Ventas por hora (S/)'],
         [[c['cajero'], c['sede'], c['turno'], c['total'], c['ventas'], c['minutos'],
           c['ticket_promedio'], c['horas'], c['ventas_por_hora']] for c in r['por_cajero']]),
        ('Ventas por sede', ['Sede', 'Total (S/)', 'Ventas', 'Minutos'],
         [[s['sede'], s['total'], s['ventas'], s['minutos']] for s in r['por_sede']]),
        ('Ventas por tipo de atracción', ['Tipo', 'Total (S/)', 'Alquileres', 'Minutos'],
         [[t['tipo'], t['total'], t['alquileres'], t['minutos']] for t in r['por_tipo']]),
        ('Ventas por método de pago', ['Método', 'Total (S/)', 'Ventas'],
         [[m['metodo'], m['total'], m['ventas']] for m in r['por_metodo']]),
        ('Ventas por día', ['Fecha', 'Total (S/)'], [[d['fecha'], d['total']] for d in r['por_dia']]),
    ]


def _titulo(r, sede_nombre):
    return f'Reporte de ventas · {r["etiqueta"]} · {sede_nombre or "Todas las sedes"}'


def _libro_excel(tablas, titulo):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, encabezados, filas in tablas:
        hoja = libro.create_sheet(nombre[:31])
        hoja.append([titulo])
        hoja['A1'].font = Font(bold=True, size=12)
        hoja.append([])
        hoja.append(encabezados)
        for celda in hoja[3]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='006C4D')
            celda.alignment = Alignment(horizontal='center', wrap_text=True)
        for fila in filas:
            hoja.append(fila)
        for col in hoja.columns:
            ancho = max((len(str(c.value)) for c in col[2:] if c.value is not None), default=10)
            hoja.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 12), 40)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _libro_pdf(tablas, titulo, max_filas=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    salida = io.BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                            topMargin=24, bottomMargin=24, title='Reporte Dino Park')
    estilos = getSampleStyleSheet()
    elementos = [Paragraph('Dino Park', estilos['Title']), Paragraph(titulo, estilos['Heading3'])]
    for nombre, encabezados, filas in tablas:
        if max_filas and len(filas) > max_filas:
            filas = filas[:max_filas]
        elementos += [Spacer(1, 10), Paragraph(nombre, estilos['Heading4'])]
        datos = [encabezados] + [[str(c) for c in f] for f in filas] if filas else [encabezados, ['Sin datos']]
        tabla = Table(datos, repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#006C4D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E7FFF1')]),
        ]))
        elementos.append(tabla)
    doc.build(elementos)
    return salida.getvalue()


def exportar_excel(r, sede_nombre=None):
    return _libro_excel(_tablas(r), _titulo(r, sede_nombre))


def exportar_pdf(r, sede_nombre=None):
    return _libro_pdf(_tablas(r), _titulo(r, sede_nombre), max_filas=31)


# ------------------------------------------------- contadores de las atracciones
def _tablas_contadores(c):
    u = lambda t: 'min' if t == 'minutos' else 'enc.'
    historial = [[f['fecha'].strftime('%d/%m/%Y %H:%M'), f['sede'], f['cajero'], f['atraccion'],
                  f['inicial'], '' if f['final'] is None else f['final'], '' if f['uso_real'] is None else f'{f["uso_real"]} {u(f["tipo_contador"])}',
                  '' if f['sistema'] is None else f['sistema'], '' if f['vendido'] is None else f['vendido'],
                  '' if f['dif_sistema'] is None else f['dif_sistema'], '' if f['dif_vendido'] is None else f['dif_vendido'],
                  'ALERTA' if f['alerta'] else 'OK', f['alertas'], f['observacion']] for f in c['filas']]
    resumen = [[r['sede'], r['atraccion'], r['tipo_contador'], r['cierres'], r['uso_real'], r['sistema'], r['vendido'], r['alertas']]
               for r in c['resumen']]
    return [
        ('Resumen por atraccion', ['Sede', 'Atraccion', 'Contador', 'Cierres', 'Uso real', 'Sistema', 'Vendido', 'Alertas'], resumen),
        ('Historial de lecturas', ['Cierre', 'Sede', 'Cajero', 'Atraccion', 'Lectura inicial', 'Lectura final', 'Uso real',
                                   'Contador sistema', 'Vendido', 'Dif. sistema', 'Dif. vendido', 'Estado', 'Alertas',
                                   'Observaciones'], historial),
    ]


def exportar_contadores_excel(c, titulo):
    return _libro_excel(_tablas_contadores(c), titulo)


def exportar_contadores_pdf(c, titulo):
    return _libro_pdf(_tablas_contadores(c), titulo, max_filas=200)
