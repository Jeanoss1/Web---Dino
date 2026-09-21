"""Turnos laborales de los cajeros y su asignación semanal."""
import csv
import io
from datetime import date, datetime, timedelta

from sqlalchemy import func

from app.database import db
from app.models.Cajero import Cajero
from app.models.Turno import AsignacionTurno, Turno
from app.models.Venta import Venta
from app.services.errores import ErrorNegocio

DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
_LIBRE = {'', 'libre', 'l', '-', 'descanso', 'x'}


def lunes_de(dia):
    return dia - timedelta(days=dia.weekday())


def dias_de_semana(lunes):
    return [lunes + timedelta(days=i) for i in range(7)]


def turnos_activos():
    return Turno.query.filter_by(estado='activo').order_by(Turno.tipo, Turno.hora_inicio).all()


def cajeros_activos(id_sucursal=None):
    consulta = Cajero.query.filter_by(estado='activo')
    if id_sucursal:
        consulta = consulta.filter_by(id_sucursal=id_sucursal)
    return sorted(consulta.all(), key=lambda c: (c.sucursal.nombre, c.tipo_turno, c.nombre))


def asignaciones_de(lunes, id_sucursal=None):
    """{(id_cajero, fecha): Turno} de la semana."""
    consulta = AsignacionTurno.query.filter(
        AsignacionTurno.fecha >= lunes, AsignacionTurno.fecha < lunes + timedelta(days=7)
    )
    if id_sucursal:
        consulta = consulta.join(Cajero).filter(Cajero.id_sucursal == id_sucursal)
    return {(a.id_cajero, a.fecha): a.turno for a in consulta.all()}


def _asignar(cajero, dia, turno):
    """Crea, cambia o quita (turno=None) la asignación de un cajero en un día."""
    actual = AsignacionTurno.query.filter_by(id_cajero=cajero.id_cajero, fecha=dia).first()
    if turno is None:
        if actual:
            db.session.delete(actual)
        return
    if turno.tipo != cajero.tipo_turno:
        raise ErrorNegocio(
            f'{cajero.nombre} es de turno {cajero.tipo_turno}; no se le puede asignar «{turno.nombre}».'
        )
    if actual:
        actual.id_turno = turno.id_turno
    else:
        db.session.add(AsignacionTurno(id_cajero=cajero.id_cajero, id_turno=turno.id_turno, fecha=dia))


def guardar_semana(lunes, seleccion):
    """seleccion = {(id_cajero, fecha): id_turno o ''}"""
    turnos = {t.id_turno: t for t in Turno.query.all()}
    cajeros = {c.id_cajero: c for c in Cajero.query.all()}
    try:
        for (id_cajero, dia), id_turno in seleccion.items():
            cajero = cajeros.get(id_cajero)
            if cajero is None or not lunes <= dia < lunes + timedelta(days=7):
                raise ErrorNegocio('Asignación fuera de la semana o cajero desconocido.')
            turno = None
            if id_turno not in ('', None):
                turno = turnos.get(int(id_turno))
                if turno is None:
                    raise ErrorNegocio('Turno desconocido.')
            _asignar(cajero, dia, turno)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise


def copiar_semana(lunes_origen, lunes_destino, rotar=True, id_sucursal=None):
    """Repite la semana de origen en el destino. Con `rotar`, cada cajero pasa al
    siguiente turno de su mismo tipo (p. ej. mañana -> tarde)."""
    origen = asignaciones_de(lunes_origen, id_sucursal)
    if not origen:
        raise ErrorNegocio('La semana de origen no tiene asignaciones.')
    por_tipo = {}
    for t in turnos_activos():
        por_tipo.setdefault(t.tipo, []).append(t)
    desplazamiento = (lunes_destino - lunes_origen).days
    cajeros = {c.id_cajero: c for c in Cajero.query.all()}
    try:
        for (id_cajero, dia), turno in origen.items():
            opciones = por_tipo.get(turno.tipo, [turno])
            nuevo = turno
            if rotar and turno in opciones:
                nuevo = opciones[(opciones.index(turno) + 1) % len(opciones)]
            _asignar(cajeros[id_cajero], dia + timedelta(days=desplazamiento), nuevo)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return len(origen)


def _filas_archivo(archivo):
    nombre = (archivo.filename or '').lower()
    contenido = archivo.read()
    if nombre.endswith('.xlsx'):
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
        return [[('' if c is None else str(c).strip()) for c in fila] for fila in hoja.iter_rows(values_only=True)]
    if nombre.endswith('.csv'):
        texto = contenido.decode('utf-8-sig', errors='replace')
        delimitador = ';' if texto.count(';') > texto.count(',') else ','
        return [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]
    raise ErrorNegocio('Sube un archivo .csv o .xlsx.')


def importar_horario(archivo, lunes):
    """Formato: una fila por cajero -> correo, Lunes..Domingo (nombre del turno o «libre»)."""
    filas = _filas_archivo(archivo)
    turnos = {t.nombre.lower(): t for t in Turno.query.all()}
    usuarios = {c.usuario.correo: c for c in Cajero.query.all()}
    cargadas = 0
    try:
        for fila in filas:
            if not fila or '@' not in (fila[0] or ''):
                continue                         # encabezado o fila vacía
            cajero = usuarios.get(fila[0].strip().lower())
            if cajero is None:
                raise ErrorNegocio(f'No existe un cajero con el correo {fila[0]}.')
            for i, dia in enumerate(dias_de_semana(lunes)):
                celda = (fila[i + 1] if len(fila) > i + 1 else '').strip()
                if celda.lower() in _LIBRE:
                    _asignar(cajero, dia, None)
                    continue
                turno = turnos.get(celda.lower())
                if turno is None:
                    raise ErrorNegocio(f'Turno «{celda}» desconocido (fila de {fila[0]}).')
                _asignar(cajero, dia, turno)
                cargadas += 1
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    if not cargadas:
        raise ErrorNegocio('El archivo no tenía asignaciones válidas.')
    return cargadas


def plantilla_csv(lunes):
    salida = io.StringIO()
    escritor = csv.writer(salida)
    escritor.writerow(['correo'] + [f'{d} {(lunes + timedelta(days=i)):%d/%m}' for i, d in enumerate(DIAS)])
    ejemplo = turnos_activos()
    nombre = ejemplo[0].nombre if ejemplo else 'Completo mañana'
    for c in cajeros_activos():
        escritor.writerow([c.usuario.correo] + [nombre] * 6 + ['libre'])
    return '﻿' + salida.getvalue()


def horas_trabajadas(cajero, desde, hasta):
    """Horas según los turnos asignados (solo hasta hoy). Si el cajero no tiene
    asignaciones en el rango, se estima: días con ventas x horas de su turno."""
    tope = min(hasta, date.today())
    if tope < desde:
        return 0.0, 'turnos'
    asignados = (
        AsignacionTurno.query.filter(
            AsignacionTurno.id_cajero == cajero.id_cajero,
            AsignacionTurno.fecha >= desde, AsignacionTurno.fecha <= tope)
        .all()
    )
    if asignados:
        return float(sum(a.turno.horas for a in asignados)), 'turnos'
    fechas = {
        f.date() if isinstance(f, datetime) else f
        for (f,) in db.session.query(Venta.fecha).filter(
            Venta.id_usuario == cajero.id_usuario, Venta.estado == 'pagada',
            Venta.fecha >= datetime.combine(desde, datetime.min.time()),
            Venta.fecha < datetime.combine(tope + timedelta(days=1), datetime.min.time()))
    }
    return float(len(fechas) * cajero.horas_turno), 'estimado'
